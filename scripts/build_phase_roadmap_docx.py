from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path("deliverables/planning-documents/CitySync_Phases_1_to_8_Workstreams.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        data = kwargs[edge]
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        for key, value in data.items():
            element.set(qn(f"w:{key}"), str(value))


def set_table_geometry(table, widths, indent=TABLE_INDENT):
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_paragraph_border(paragraph, bottom_color=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "7")
    bottom.set(qn("w:color"), bottom_color)
    p_bdr.append(bottom)


def page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_num_pr(paragraph, num_id, ilvl=0):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl_el)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def add_list_numbering(document, abstract_id, num_id, kind):
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(text)
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "left")
    level.append(justify)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(rfonts)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    def make_style(name, base="Normal"):
        if name in styles:
            return styles[name]
        return styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=False)

    title = make_style("Roadmap Title")
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.line_spacing = 1.0

    subtitle = make_style("Roadmap Subtitle")
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.15

    kicker = make_style("Roadmap Kicker")
    kicker.font.name = "Calibri"
    kicker._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    kicker._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    kicker.font.size = Pt(10)
    kicker.font.bold = True
    kicker.font.color.rgb = RGBColor.from_string(GOLD)
    kicker.paragraph_format.space_after = Pt(9)
    kicker.paragraph_format.line_spacing = 1.0

    lead = make_style("Roadmap Lead")
    lead.font.name = "Calibri"
    lead._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    lead._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    lead.font.size = Pt(12)
    lead.font.color.rgb = RGBColor.from_string(INK)
    lead.paragraph_format.space_after = Pt(10)
    lead.paragraph_format.line_spacing = 1.25

    small = make_style("Roadmap Small")
    small.font.name = "Calibri"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    small.font.size = Pt(9.5)
    small.font.color.rgb = RGBColor.from_string(MUTED)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.15

    table_text = make_style("Roadmap Table Text")
    table_text.font.name = "Calibri"
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_text.font.size = Pt(9.5)
    table_text.font.color.rgb = RGBColor.from_string(INK)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.12

    table_header = make_style("Roadmap Table Header")
    table_header.font.name = "Calibri"
    table_header._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_header._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_header.font.size = Pt(9.5)
    table_header.font.bold = True
    table_header.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    table_header.paragraph_format.space_after = Pt(0)
    table_header.paragraph_format.line_spacing = 1.0


def setup_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("CITY/SYNC  |  ROADMAP WORKSTREAMS")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    set_paragraph_border(p, bottom_color="D7DFE8")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("City/Sync Planning Guide  •  ")
    set_run_font(r, size=8.5, color=MUTED)
    page_field(p)


def add_para(doc, text="", style="Normal", align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_with_next = keep
    if text:
        p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_num_pr(p, 902)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_num_pr(p, 903)
    p.add_run(text)
    return p


def add_callout(doc, title, text, trailing_spacer=True):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_border(
        cell,
        top={"val": "single", "sz": "4", "color": "D7DFE8"},
        bottom={"val": "single", "sz": "4", "color": "D7DFE8"},
        left={"val": "single", "sz": "12", "color": BLUE},
        right={"val": "single", "sz": "4", "color": "D7DFE8"},
    )
    p = cell.paragraphs[0]
    p.style = doc.styles["Roadmap Table Text"]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + ": ")
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    if trailing_spacer:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_geometry(table, widths)
    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "4", "color": "B9C8D9"},
            bottom={"val": "single", "sz": "4", "color": "B9C8D9"},
            left={"val": "single", "sz": "4", "color": "B9C8D9"},
            right={"val": "single", "sz": "4", "color": "B9C8D9"},
        )
        p = cell.paragraphs[0]
        p.style = doc.styles["Roadmap Table Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(text)
    set_repeat_table_header(table.rows[0])
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for col_i, (cell, text) in enumerate(zip(cells, row)):
            if row_i % 2 == 1:
                set_cell_shading(cell, "FAFBFC")
            set_cell_border(
                cell,
                top={"val": "single", "sz": "2", "color": "D9E0E8"},
                bottom={"val": "single", "sz": "2", "color": "D9E0E8"},
                left={"val": "single", "sz": "2", "color": "D9E0E8"},
                right={"val": "single", "sz": "2", "color": "D9E0E8"},
            )
            p = cell.paragraphs[0]
            p.style = doc.styles["Roadmap Table Text"]
            if col_i == 0:
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            else:
                p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_phase(doc, phase, title, purpose, dependencies, workstreams, gate, measures, start_new_page=False):
    if start_new_page:
        doc.add_page_break()
    add_para(doc, f"PHASE {phase}", style="Roadmap Kicker")
    add_para(doc, title, style="Heading 1")
    add_para(doc, purpose, style="Roadmap Lead")
    add_callout(doc, "Entry condition", dependencies)
    add_para(doc, "Workstreams", style="Heading 2")
    add_matrix(
        doc,
        ["Workstream", "Achievable tasks", "Deliverable / acceptance check"],
        workstreams,
        [1850, 4700, 2810],
    )
    add_callout(doc, "Phase gate", gate)
    add_para(doc, "Core measures", style="Heading 2")
    for measure in measures:
        add_bullet(doc, measure)


def build_document():
    doc = Document()
    configure_styles(doc)
    setup_page(doc)
    add_list_numbering(doc, 902, 902, "bullet")
    add_list_numbering(doc, 903, 903, "decimal")

    # Editorial-cover header pattern: a generous, centered opening block.
    add_para(doc, "CITY/SYNC PLANNING GUIDE", style="Roadmap Kicker", align=WD_ALIGN_PARAGRAPH.CENTER)
    title = add_para(doc, "Roadmap Workstreams", style="Roadmap Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    subtitle = add_para(doc, "Phases 1–8: From nonprofit product to federated civic infrastructure", style="Roadmap Subtitle", align=WD_ALIGN_PARAGRAPH.CENTER)
    subtitle.paragraph_format.space_after = Pt(24)
    add_para(doc, "August 2026", style="Roadmap Small", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "A sequencing guide built from the City/Sync backend target architecture and current product roadmap.", style="Roadmap Small", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    add_callout(
        doc,
        "How to use this guide",
        "Each phase is organized as workstreams with concrete outputs and an explicit readiness gate. Move forward only when the prior phase has demonstrated real value, safe operations, and accountable governance—not merely when software features are complete.",
    )
    add_para(doc, "Roadmap at a glance", style="Heading 1")
    add_matrix(
        doc,
        ["Phase", "Strategic purpose", "Proof before advancing"],
        [
            ("1", "Become indispensable to volunteer organizations.", "Organizations save time and retain volunteers through daily use."),
            ("2", "Form a functioning local City/Sync network.", "Organizations and participants share a trusted local experience."),
            ("3", "Prove coordinated civic action at city scale.", "One visible Mass Coordination Event produces a verified public outcome."),
            ("4", "Build the city policy and accepted-event protocol.", "Organizations and city domains exchange verifiable, privacy-safe claims."),
            ("5", "Integrate public institutions without replacing their systems.", "One agency contributes accepted records and obtains a real reporting benefit."),
            ("6", "Introduce governed local recognition and Civic Credits.", "A funded, reconciled, equity-reviewed local program operates safely."),
            ("7", "Enable independent operation and federation.", "Organizations and cities can export, hand off, and independently verify continuity."),
            ("8", "Evaluate shared execution only if it is necessary.", "Anchored databases demonstrably no longer satisfy network needs."),
        ],
        [720, 4200, 4440],
    )
    add_para(doc, "Planning assumptions", style="Heading 1")
    for item in [
        "Phase 0 establishes launch readiness: one primary V1 interface, staging and production hygiene, testing, security, privacy, support, and a closed real-world pilot.",
        "The present application is database-first, event-journaled, city-partitioned, and chain-ready. It is not yet a fully federated organization-node and city-acceptance protocol.",
        "Civic Credits, public-chain anchoring, and shared execution remain optional modules. They are introduced only after their governance, funding, and operational conditions are met.",
        "City/Sync should lead with practical institutional value—volunteer operations, reporting, coordination, and evidence lineage—not with blockchain terminology.",
    ]:
        add_bullet(doc, item)

    add_phase(
        doc,
        "1",
        "Nonprofit Optimization",
        "Make City/Sync the practical operating system that volunteer organizations choose because it saves staff time, improves volunteer experience, and produces usable evidence of program activity.",
        "Phase 0 launch gate has passed: the selected V1 interface is stable, the pilot operating model is documented, and real-user feedback has identified the highest-friction workflows.",
        [
            ("Daily volunteer operations", "Refine opportunity publishing, recurring schedules, capacity controls, waiver flow, attendance, completion, verification, cancellation, and no-show handling. Reduce the number of steps required for staff and participants.", "A complete volunteer lifecycle works for a live organization without off-platform spreadsheets."),
            ("Organization workspace", "Improve roster search, grouping, messaging, role-based access, delegated authority, public profile, locations, onboarding sessions, and reusable opportunity templates.", "Organization owners can configure their account, invite staff, and run daily operations independently."),
            ("Participant retention", "Improve discovery, saved opportunities, service history, notifications, calendar views, accessibility, mobile web behavior, and clear city-participation status.", "Participants can find, claim, complete, and understand the value of local participation without training."),
            ("Compliance and reporting", "Map common nonprofit reporting needs; add report definitions, exports, waiver history, participant-hours summaries, organization activity reports, and traceable filters.", "An organization can produce a useful operational or funder report from City/Sync data in minutes."),
            ("Integrations and migration", "Prioritize calendar, email, spreadsheet import/export, public-website embed, and basic API/webhook planning. Create a practical migration checklist for an organization moving from spreadsheets or another tool.", "The product can enter an organization’s existing workflow instead of demanding an all-at-once replacement."),
            ("Success and support", "Create organization health metrics, in-product guidance, onboarding templates, feedback interviews, support classification, and a recurring product-review cadence.", "Retention, feature use, support demand, and product gaps are measured per organization."),
        ],
        "At least 10 organizations use the platform for real volunteer operations; organizations report a clear time or coordination benefit; the core loop is reliable enough that manual staff intervention is exceptional rather than routine.",
        [
            "Organization activation, first opportunity published, first volunteer completion, and 30/90-day organization retention.",
            "Volunteer claim-to-completion rate, verification turnaround time, no-show rate, and repeat participation.",
            "Time spent producing a recurring report compared with each organization’s prior workflow.",
        ],
    )

    add_phase(
        doc,
        "2",
        "City Network Formation",
        "Move from isolated organization accounts to a coherent local network where residents discover trusted opportunities, understand membership, and participate across recognized institutions.",
        "Phase 1 has demonstrated organization-level product value, and one local city operator or chapter has accepted responsibility for operating the first city network.",
        [
            ("City operator and governance", "Name the city operator, policy authority, privacy contact, participant-support contact, and organization-recognition process. Define the local charter and decision rights.", "A public city-network operating charter identifies who runs the network and how consequential decisions are made."),
            ("City-domain setup", "Provision the city domain, configure city-specific policies, organization status, local administrator roles, data boundaries, backups, and incident contacts.", "The city has an isolated operational domain with named owners and documented recovery responsibilities."),
            ("Organization recognition", "Create application, due-diligence, approval, suspension, correction, renewal, and public-directory workflows for local organizations.", "Residents can distinguish recognized organizations from unreviewed listings."),
            ("City membership", "Define New Participant and City Member states, onboarding requirements, neighboring-city participation rules, no-show recovery, and participant correction paths.", "Membership status is understandable, consistently applied, and visible to participants."),
            ("Shared local experience", "Operate city discovery, organization directory, opportunity calendar, MyCity Feed, notifications, local service history, and public city-impact views.", "A participant experiences one coherent local network rather than disconnected organization tools."),
            ("Local chapter capacity", "Recruit and train local facilitators, set support routines, build partner relationships, and document chapter playbooks for a second city.", "The city network can operate weekly without relying on one central founder for every decision."),
        ],
        "One city network operates with recognized organizations, an active participant pathway, a published operating charter, and a local operator capable of handling ordinary support and organization approval.",
        [
            "Recognized organizations, active participants, city-member activation rate, and cross-organization participation.",
            "City-directory use, opportunity discovery-to-claim conversion, and local feed/notification engagement.",
            "Time required to approve, onboard, and support an organization within the city network.",
        ],
    )

    add_phase(
        doc,
        "3",
        "Mass Coordination Events",
        "Demonstrate that the city network can organize a visible, bounded public effort that no single organization could deliver as effectively on its own.",
        "Phase 2 has established a functioning local network, and an appropriate public challenge has been selected with local partners, clear safety boundaries, and accountable sponsorship.",
        [
            ("Challenge selection", "Choose one bounded public problem with a clear geographic scope, measurable outcome, partner agencies, risk review, and public value. Avoid high-risk tasks in the first event.", "A published challenge brief explains the need, intended outcome, limits, partners, and success measures."),
            ("Public plan and task design", "Break the objective into discrete work units, define team size, location, skills, supplies, safety, permissions, accessibility, time windows, and completion criteria.", "Every task is understandable, claimable, and connected to a visible portion of the whole plan."),
            ("Teams and briefings", "Build team-formation workflows, random team option, team leads, briefing materials, meeting points, escalation contacts, and follow-up instructions.", "Participants can join or form a team and receive the information needed to act safely."),
            ("Equipment and permissions", "Create equipment-team roles, inventory/checkout procedures, site permissions, insurance/waiver review, transport plan, and return verification.", "The event has a documented logistics plan; equipment and permissions do not become hidden blockers."),
            ("Verification and after-action record", "Recruit qualified verifiers, establish sampling and correction rules, capture permitted evidence, build a public progress view, and assign after-action documentation tasks.", "Completed work has an accountable verification path and the event produces a reusable after-action report."),
            ("Community learning", "Run a participant debrief, partner review, public results release, cost/benefit review, and policy update process before repeating the model.", "The city network publishes what worked, what did not, and what will change next time."),
        ],
        "One Mass Coordination Event is completed safely, has a measurable public outcome, and produces an independently reviewable internal record and public after-action report.",
        [
            "Participants and teams activated; task fill, completion, verification, correction, and abandonment rates.",
            "Outcome units completed against plan; logistic incidents; accessibility and safety issues; partner satisfaction.",
            "Time from public challenge to organized action and the amount of coordination performed across institutions.",
        ],
    )

    add_phase(
        doc,
        "4",
        "City Ledger and Policy Protocol",
        "Evolve the current event-journal and city-outbox foundation into a privacy-safe, signed protocol through which organization claims are accepted into city-level civic state.",
        "Phases 1–3 have provided real workflow evidence. The first city has named policy authority, technical stewardship, privacy ownership, and a narrow initial event type to standardize.",
        [
            ("Control-plane minimization", "Separate platform authentication, routing, discovery, and software catalogs from substantive organization and city records. Document temporary duplicates and their migration path.", "The control plane contains only the minimum information necessary to route and operate the network."),
            ("Organization-domain isolation", "Provision isolated organization domains; move projections, source journals, outboxes, receipt state, and evidence metadata into each organization boundary. Deliver export and restore for managed nodes.", "A pilot organization can operate and recover without sharing its substantive database with another organization."),
            ("Envelope and signature v1", "Publish canonical source-envelope, error, receipt, and hash-suite schemas. Add source sequence, authority reference, policy reference, payload digest, privacy class, key lifecycle, and signatures.", "Two independent implementations can reproduce and verify identical source hashes and signed envelopes from fixtures."),
            ("City intake and acceptance", "Build validation for membership, supported schemas, signatures, sequence gaps, duplicate delivery, authority, policy, privacy, evidence status, quarantine, and signed receipts.", "Invalid events cannot change city projections; retries have one accepted effect; forks stop safely."),
            ("Policy and authority registry", "Implement policy-package lifecycle, machine-readable rules, tests, authority delegations, city trust registries, revocation, and explanation templates.", "A disclosed accepted event can be traced to the authority and policy version that governed it."),
            ("Proofs, reporting, and anchors", "Create evidence descriptors, report manifests, domain-separated Merkle batches, proof bundles, independent verifier tooling, and privacy review for public metadata.", "A permitted verifier can reproduce a report lineage or disclosed event proof without accessing unrelated private data."),
        ],
        "A real organization submits signed, minimal completion claims; the city validates and accepts them under a published policy; participants can see status and correction paths; independent verification succeeds on the defined proof scope.",
        [
            "Source-to-city acceptance latency, outbox age, duplicate rate, gap/fork rate, quarantine reasons, and reconciliation failures.",
            "Policy/authority resolution success, proof-verification success, correction time, and participant record-understanding feedback.",
            "City-isolation tests, organization export/restore tests, and private-data disclosure exceptions.",
        ],
    )

    add_phase(
        doc,
        "5",
        "Public Institution Integration",
        "Allow local government agencies to obtain shared coordination and reporting value while keeping their own operational systems, private records, and institutional authority.",
        "Phase 4 has established a stable accepted-event protocol, a narrow tested policy profile, and a local governance body capable of reviewing agency participation and data boundaries.",
        [
            ("Institutional offer", "Package City/Sync as a Verifiable Program Operations Pilot. Define scope, service levels, security posture, records ownership, exit rights, implementation support, and procurement-friendly language.", "An agency can evaluate City/Sync as practical civic infrastructure rather than a speculative technology project."),
            ("Integration profiles", "Publish API, webhook, database/change-log connector, scheduled signed-import, and manual-attestation profiles. Define assurance levels and which effects each profile may trigger.", "Agencies can integrate without replacing their system, while the city can distinguish high-assurance from lower-assurance submissions."),
            ("Policy and data-sharing templates", "Create templates for volunteer completion, training credential, public-project milestone, and aggregate reporting. Include purpose, minimum fields, retention, correction, recourse, privacy, and evidence rules.", "Every agency integration starts from an approved policy and data-minimization template rather than custom ad hoc sharing."),
            ("First agency pilot", "Choose a low-risk program such as parks, libraries, public works, emergency preparedness, or community stewardship. Map fields, register authority keys, run sandbox fixtures, then conduct a limited live pilot.", "One agency produces accepted events from its existing system and retains its private source records."),
            ("Evidence-derived reporting", "Build the first agency report manifest and public aggregate result. Link outcomes to accepted events, policy versions, exclusions, corrections, and reviewer authority.", "The agency can produce a useful program or funder report faster and with clearer lineage than before."),
            ("Adoption and accountability", "Create agency onboarding, public communication, accessibility review, privacy/civil-rights review, staff training, support escalation, and annual participation review.", "The agency can explain what City/Sync does, what it does not do, and how residents can seek correction or remedy."),
        ],
        "One local government agency completes a real, low-risk program integration; its records remain in its own system; the city accepts only approved minimal claims; and the agency obtains a demonstrated coordination or reporting benefit.",
        [
            "Integration time, connector reliability, acceptance/rejection outcomes, reconciliation effort, and agency support demand.",
            "Reporting preparation time, traceability of reported outcomes, participant correction requests, and privacy incidents.",
            "Agency partner satisfaction and the proportion of program work coordinated without duplicating private records into a central warehouse.",
        ],
    )

    add_phase(
        doc,
        "6",
        "Governed Civic Credits",
        "Introduce a city-scoped recognition and redemption program only after local policy, funding, verification, equity, and operational controls are mature.",
        "Phase 5 has demonstrated accepted-event reliability and local institutions have endorsed a specific, funded local-recognition purpose. Legal, fiscal, labor, accessibility, and equity review is complete.",
        [
            ("Program charter", "Define the local purpose, eligibility, exclusions, issuance caps, non-transferability, expiration if any, appeal rights, conflict rules, and relationship to paid labor and essential services.", "A public Civic Credit charter makes clear that credits are local recognition units, not investment assets or a replacement for public benefits."),
            ("Funding and benefits", "Secure funded benefit partners, benefit inventory, budget owner, redemptions policy, fulfillment evidence, reserve rules, and failure/closure plan.", "Every redemption promise has a named provider, funding source, and accountable fulfillment path."),
            ("Credit-journal hardening", "Make the city credit journal authoritative. Implement idempotent mint, adjustment, reservation, release, burn, issuer caps, balance protection, reconciliation, and pause controls.", "Every balance change traces to an eligible accepted event and a specific policy rule; no balance can become negative."),
            ("Participant experience", "Design plain-language wallet, receipt, redemption, dispute, privacy, and accessibility flows. Avoid wallet keys, speculative language, and coercive participation incentives.", "Participants understand what they earned, why, what can be redeemed, and how to correct a mistake."),
            ("Equity and anti-gaming", "Test for exclusion, labor substitution, geographic inequity, fraud, collusion, verifier conflicts, and uneven access to benefits. Establish independent review and corrective authority.", "The pilot has documented safeguards and can pause safely when anomalies or equity risks appear."),
            ("Controlled launch", "Start with a narrow city program, small cap, limited benefits, clear communications, reconciliation cadence, and public aggregate reporting.", "A local credit pilot completes multiple issuance and redemption cycles without unreconciled balances or unresolved participant harm."),
        ],
        "A city-scoped Civic Credit program operates with funded benefits, reconciled balances, correction/appeal procedures, equity safeguards, and no dependence on a transferable token or cryptocurrency wallet.",
        [
            "Issuance and redemption reconciliation, outstanding reservations, fulfillment time, dispute rate, and pause incidents.",
            "Participation and benefit access across neighborhoods and groups; evidence of labor displacement, gaming, or coercion.",
            "Administrative cost per issued/redeemed credit and funded-benefit sustainability.",
        ],
    )

    add_phase(
        doc,
        "7",
        "Sovereign Nodes and Federation",
        "Make City/Sync genuinely replaceable: organizations and cities can operate compatible domains, retain control of records and keys, and exchange proofs without merging into one global ledger.",
        "Phase 6 has proven operational value. The protocol schemas, proofs, exports, and city governance model are stable enough for third parties to operate without relying on informal institutional knowledge.",
        [
            ("Open protocol package", "Publish schemas, canonicalization rules, signature profiles, conformance fixtures, error codes, compatibility profiles, verifier instructions, and version-support policy.", "An external implementation can build and test a compatible organization or city domain from public materials."),
            ("Managed-node handoff", "Deliver export, encrypted transfer, integrity manifest, restore, replay, routing switch, key transition, city receipt reconciliation, and rollback procedures for an organization leaving managed hosting.", "A managed organization moves to another operator without losing history or creating duplicate city effects."),
            ("Sovereign organization pilot", "Select one capable organization or local operator to run the reference node. Test keys, backups, policy cache, ordered outbox, receipt handling, incident response, and support boundaries.", "A sovereign organization node submits valid events and remains compatible with city intake."),
            ("Sovereign city capabilities", "Package city-domain deployment, operator console, key ceremony, backup/restore, report/proof archive, monitoring, and stewardship-transfer runbooks.", "A city can run or migrate its own domain with documented recovery and accountability controls."),
            ("Federated credentials and proofs", "Pilot one portable service, training, or qualification credential between two cities. Require participant consent and receiving-city policy acceptance; do not move balances or raw records.", "A second city verifies a permitted proof and records its own local acceptance without hidden cross-city correlation."),
            ("Protocol governance", "Establish protocol-change proposals, security coordination, privacy review, conformance certification, release support windows, and multi-stakeholder decision rights.", "A breaking change or security issue can be managed without unilateral, opaque platform control."),
        ],
        "At least one organization and one city can replace the managed operator or operate independently while preserving records, verification continuity, privacy boundaries, and protocol compatibility.",
        [
            "Export/restore success, operator-handoff duration, independent verifier success, node uptime, outbox reconciliation, and recovery-drill results.",
            "Credential portability acceptance/rejection outcomes and cross-city privacy/correlation incidents.",
            "Protocol conformance adoption, operator support burden, governance participation, and version-migration success.",
        ],
    )

    add_phase(
        doc,
        "8",
        "Conditional Shared Execution",
        "Evaluate shared on-chain registries, a dedicated execution layer, or L2/L3-style settlement only if the federated, anchored database architecture has proven insufficient for a demonstrated public need.",
        "Phase 7 has multiple independent operators, stable governance, sustainable funding, audited operational controls, and documented evidence that anchoring alone cannot meet a required interoperability or settlement need.",
        [
            ("Decision threshold", "Publish the specific unmet requirement: cross-city settlement, shared registry, finality, programmatic escrow, interoperability, or another capability. Compare non-chain alternatives and cost them honestly.", "A public architecture decision explains why shared execution is needed and what simpler approach failed to satisfy."),
            ("Architecture selection", "Evaluate execution environment, settlement chain, data availability, finality, operator model, chain abstraction, privacy, exit, cost, jurisdiction, and independent-verification implications.", "A chosen design meets documented security, governance, cost, and exit criteria—not merely ecosystem preference."),
            ("Prototype and sandbox", "Build a non-production proof of concept for the smallest necessary shared registry or settlement module. Keep personal data and private evidence off chain; test chain outages and fallback behavior.", "The prototype demonstrates a real capability unavailable through anchored city domains and does not create a new privacy or operational dependency."),
            ("Security and governance", "Commission threat modeling, contract audits, key/threshold design, sequencer or operator governance, incident response, upgrade controls, emergency pause, and public disclosure processes.", "Independent review confirms that operational and governance risks are understood before funds or public rights depend on contracts."),
            ("Economic sustainability", "Model gas sponsorship, institutional subscriptions, settlement costs, grants, reserve policy, operator expenses, and long-term maintenance. Keep residents gasless without pretending infrastructure is free.", "The shared layer has a funded operating model, clear payer, cost ceiling, and contingency plan."),
            ("Limited production use", "Deploy only the narrowly justified module with slow rollout, caps, monitoring, proof of finality, adapter fallback, participant communications, and independent oversight.", "The system can pause or replace shared execution without stopping ordinary local organization and city operations."),
        ],
        "Shared execution is adopted only for a narrowly defined, audited, funded, governed function that produces measurable public value beyond the anchored federated architecture.",
        [
            "Demonstrated reduction in reconciliation or coordination burden; settlement/finality cost; availability; audit findings; and fallback success.",
            "Institutional funding coverage, governance participation, incident count, and operator concentration risk.",
            "Evidence that residents remain gasless, private data remains private, and local autonomy is preserved.",
        ],
    )

    add_para(doc, "Cross-Phase Workstreams", style="Heading 1")
    add_para(doc, "These workstreams do not belong to only one phase. They must mature alongside the product and protocol.", style="Roadmap Lead")
    add_matrix(
        doc,
        ["Workstream", "Continue throughout the roadmap"],
        [
            ("Community and chapter development", "Train local partners, build inclusive participation routes, and develop durable chapter playbooks."),
            ("Privacy, security, and civil rights", "Maintain minimization, access and retention rules, correction, accessibility, incident response, and independent review."),
            ("Funding and sustainability", "Combine grants, institutional contracts, delivery services, and aligned sponsorship without charging residents to participate."),
            ("Research and evaluation", "Measure civic capacity, administrative burden, equity, public outcomes, trust, retention, and unintended effects—not only account volume."),
            ("Mobile and accessibility", "Support mobile, low-bandwidth, multilingual, offline where appropriate, and non-digital participation routes."),
            ("Public communication", "Explain what City/Sync does and does not verify, who governs rules, and how residents seek correction or inspect results."),
        ],
        [2200, 7160],
    )
    add_para(doc, "Sequencing guardrails", style="Heading 1")
    for item in [
        "Civic Credits require funded benefits, reliable accepted events, equity safeguards, and accountable local governance.",
        "Agencies retain their databases; integrate only policy-bound minimal claims that demonstrate reporting or coordination value.",
        "Do not claim federation until export, handoff, independent verification, and operator replacement are field-tested.",
        "Build shared execution only if anchored city domains cannot meet an evidenced public need.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "Immediate next move", "Finish Phase 0 and choose one V1 interface. Then prove durable nonprofit value in Phase 1 before adding city-level complexity.", trailing_spacer=False)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "City/Sync Roadmap Workstreams: Phases 1–8"
    doc.core_properties.subject = "Implementation workstreams and readiness gates"
    doc.core_properties.author = "City/Sync"
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build_document()
